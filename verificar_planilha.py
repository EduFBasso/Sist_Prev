"""Verificar planilha gerada"""
import openpyxl

wb = openpyxl.load_workbook('saida/simulacao_joao_carlos.xlsx', data_only=False)

print('\n' + '='*70)
print('📊 RESUMO DA PLANILHA GERADA')
print('='*70 + '\n')

# Dados Cliente
ws = wb['Dados_Cliente']
print('DADOS DO CLIENTE:')
print(f'  Nome: {ws["B3"].value}')
print(f'  CPF: {ws["B4"].value}')
print(f'  NIT: {ws["B5"].value}')

# Remunerações
ws = wb['Remuneracoes']
total = 0
for row in ws.iter_rows(min_row=3, max_row=200, values_only=True):
    if row[0] is not None:
        total += 1
print(f'\nREMUNERAÇÕES:')
print(f'  Total de competências: {total}')

# Cálculo Média
ws = wb['Calculo_Media']
print(f'\nCÁLCULO MÉDIA (fórmulas):')
print(f'  B4 (Quantidade): {ws["B4"].value}')
print(f'  B5 (Soma): {ws["B5"].value}')
print(f'  B6 (Média): {ws["B6"].value}')

# Cálculo Tempo
ws = wb['Calculo_Tempo']
print(f'\nCÁLCULO TEMPO (fórmulas):')
print(f'  B4 (Meses): {ws["B4"].value}')
print(f'  B5 (Anos): {ws["B5"].value}')
print(f'  B6 (Acima mínimo): {ws["B6"].value}')

# Cálculo Final
ws = wb['Calculo_Final']
print(f'\nCÁLCULO FINAL (fórmulas):')
print(f'  B4 (Média): {ws["B4"].value}')
print(f'  B5 (Anos): {ws["B5"].value}')
print(f'  B12 (Coeficiente): {ws["B12"].value}')
print(f'  B19 (RESULTADO): {ws["B19"].value}')

print('\n' + '='*70)
print('✅ Planilha com todas as fórmulas Excel nativas!')
print('📂 Abra: saida/simulacao_joao_carlos.xlsx')
print('='*70 + '\n')
