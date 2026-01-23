"""
Gerar Planilha Simples - Teste do Coordenador

Cria planilha Excel mostrando as remunerações extraídas pelo coordenador.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from extrator.tipos.coordenador_remuneracoes import (
    extrair_remuneracoes_coordenado,
    validar_baseline_remuneracoes
)

PDF_TESTE = "cnis/CNIS_JOAO_CARLOS.pdf"
SAIDA_EXCEL = "saida/teste_coordenador.xlsx"

def main():
    print("=" * 70)
    print("📊 GERANDO PLANILHA SIMPLES - TESTE COORDENADOR")
    print("=" * 70)
    print()
    
    # Extrair remunerações
    print(f"📄 Processando: {PDF_TESTE}")
    registros = extrair_remuneracoes_coordenado(PDF_TESTE)
    print(f"✅ Extraídas: {len(registros)} remunerações")
    print()
    
    # Validar
    validacao = validar_baseline_remuneracoes(registros)
    print(f"📊 Status: {validacao['mensagem']}")
    print()
    
    # Criar planilha
    print(f"💾 Criando planilha: {SAIDA_EXCEL}")
    wb = Workbook()
    
    # Aba 1: Remunerações
    ws = wb.active
    ws.title = "Remunerações"
    
    # Header
    headers = ["Página", "Seq", "CNPJ", "Tipo", "Competência", "Remuneração", "Indicadores"]
    ws.append(headers)
    
    # Formatar header
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Dados
    for reg in registros:
        ws.append([
            reg.get('pagina', ''),
            reg.get('seq', ''),
            reg.get('cnpj', ''),
            reg.get('tipo_vinculo', ''),
            reg.get('competencia', ''),
            float(reg.get('remuneracao', 0)) if reg.get('remuneracao') else '',
            reg.get('indicadores', '')
        ])
    
    # Formatar colunas
    ws.column_dimensions['A'].width = 10  # Página
    ws.column_dimensions['B'].width = 8   # Seq
    ws.column_dimensions['C'].width = 22  # CNPJ
    ws.column_dimensions['D'].width = 12  # Tipo
    ws.column_dimensions['E'].width = 14  # Competência
    ws.column_dimensions['F'].width = 15  # Remuneração
    ws.column_dimensions['G'].width = 30  # Indicadores
    
    # Formatar valores monetários
    for row in range(2, len(registros) + 2):
        cell = ws.cell(row=row, column=6)  # Coluna F (Remuneração)
        cell.number_format = '#,##0.00'
        cell.alignment = Alignment(horizontal="right")
    
    # Aba 2: Estatísticas
    ws_stats = wb.create_sheet("Estatísticas")
    
    # Título
    ws_stats['A1'] = "ESTATÍSTICAS - COORDENADOR DE REMUNERAÇÕES"
    ws_stats['A1'].font = Font(bold=True, size=14)
    ws_stats.merge_cells('A1:D1')
    
    # Total
    ws_stats['A3'] = "Total de remunerações:"
    ws_stats['B3'] = len(registros)
    ws_stats['B3'].font = Font(bold=True, size=12)
    
    # Baseline esperado
    ws_stats['A4'] = "Baseline esperado:"
    ws_stats['B4'] = 178
    
    # Percentual
    percentual = (len(registros) / 178 * 100) if len(registros) > 0 else 0
    ws_stats['A5'] = "Percentual capturado:"
    ws_stats['B5'] = f"{percentual:.1f}%"
    ws_stats['B5'].font = Font(bold=True, color="FF0000" if percentual < 100 else "00FF00")
    
    # Por tipo
    ws_stats['A7'] = "POR TIPO:"
    ws_stats['A7'].font = Font(bold=True)
    
    row = 8
    for tipo, count in validacao['por_tipo'].items():
        ws_stats[f'A{row}'] = f"  {tipo}:"
        ws_stats[f'B{row}'] = count
        row += 1
    
    # Por sequência
    ws_stats['A11'] = "POR SEQUÊNCIA:"
    ws_stats['A11'].font = Font(bold=True)
    
    por_seq = {}
    for reg in registros:
        seq = reg.get('seq', '?')
        por_seq[seq] = por_seq.get(seq, 0) + 1
    
    row = 12
    for seq in sorted(por_seq.keys(), key=lambda x: int(x) if x.isdigit() else 999):
        ws_stats[f'A{row}'] = f"  Seq {seq}:"
        ws_stats[f'B{row}'] = por_seq[seq]
        row += 1
    
    # Larguras
    ws_stats.column_dimensions['A'].width = 30
    ws_stats.column_dimensions['B'].width = 15
    
    # Salvar
    wb.save(SAIDA_EXCEL)
    
    print(f"✅ Planilha salva!")
    print()
    print("─" * 70)
    print("📋 CONTEÚDO DA PLANILHA:")
    print("─" * 70)
    print()
    print(f"Aba 1 - Remunerações: {len(registros)} linhas")
    print(f"  Colunas: Página | Seq | CNPJ | Tipo | Competência | Remuneração | Indicadores")
    print()
    print(f"Aba 2 - Estatísticas:")
    print(f"  Total: {len(registros)}/178 ({percentual:.1f}%)")
    print(f"  Por tipo:")
    for tipo, count in validacao['por_tipo'].items():
        print(f"    {tipo}: {count}")
    print()
    print(f"  Por sequência: {len(por_seq)} diferentes")
    print()
    
    # Análise de faltantes
    print("─" * 70)
    print("⚠️  ANÁLISE DE FALTANTES:")
    print("─" * 70)
    print()
    
    # Baseline esperado por seq
    baseline_por_seq = {
        '1': 5, '2': 31, '3': 2, '4': 6, '5': 3, '6': 4, '7': 7,
        '8': 38, '9': 21, '10': 46, '11': 2, '12': 9, '13': 4
    }
    
    for seq, esperado in baseline_por_seq.items():
        capturado = por_seq.get(seq, 0)
        faltam = esperado - capturado
        
        if faltam > 0:
            print(f"Seq {seq:2s}: {capturado:2d}/{esperado:2d} (faltam {faltam:2d}) ⚠️")
        elif faltam == 0:
            print(f"Seq {seq:2s}: {capturado:2d}/{esperado:2d} ✅")
        else:
            print(f"Seq {seq:2s}: {capturado:2d}/{esperado:2d} (excesso {abs(faltam):2d}) ❓")
    
    print()
    print("=" * 70)
    print(f"📊 Arquivo: {SAIDA_EXCEL}")
    print("=" * 70)

if __name__ == "__main__":
    main()
