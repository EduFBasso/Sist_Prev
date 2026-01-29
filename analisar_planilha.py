"""Análise da estrutura da planilha inss_simulacao.xlsx"""
import openpyxl

wb = openpyxl.load_workbook('Planilhas/inss_simulacao.xlsx', data_only=True)

print("=" * 80)
print("ESTRUTURA DA PLANILHA: inss_simulacao.xlsx")
print("=" * 80)

for sheet_name in wb.sheetnames:
    print(f"\n{'='*80}")
    print(f"ABA: {sheet_name}")
    print(f"{'='*80}\n")
    
    ws = wb[sheet_name]
    
    # Primeiras 25 linhas
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True), 1):
        if any(cell is not None for cell in row):
            # Formatar linha removendo Nones do final
            row_data = list(row)
            while row_data and row_data[-1] is None:
                row_data.pop()
            print(f"L{i:02d}: {row_data}")

wb.close()
